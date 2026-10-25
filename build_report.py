from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'Kitsune_Sushi_Website_Report.docx'

def set_font(run, size=12, bold=False, italic=False):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic

def add_link(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), rel_id)
    run = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); rpr.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); rpr.append(underline)
    rfonts = OxmlElement('w:rFonts'); rfonts.set(qn('w:ascii'), 'Times New Roman'); rfonts.set(qn('w:hAnsi'), 'Times New Roman'); rpr.append(rfonts)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '24'); rpr.append(sz)
    run.append(rpr)
    t = OxmlElement('w:t'); t.text = text; run.append(t)
    hyperlink.append(run); paragraph._p.append(hyperlink)

def add_para(doc, text='', bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Inches(.25)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):]); set_font(r)
    else:
        r = p.add_run(text); set_font(r)
    return p

def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text); set_font(r, bold=True)
    return p

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.5); sec.bottom_margin = Inches(.5)
sec.left_margin = Inches(.6); sec.right_margin = Inches(.6)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1
normal.paragraph_format.space_after = Pt(2)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(5)
r = title.add_run('Kitsune Sushi Website Report'); set_font(r, size=14, bold=True)

p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
r = p.add_run('Website URL: '); set_font(r, bold=True)
add_link(p, 'http://localhost:8443', 'http://localhost:8443')
r = p.add_run(' (local development preview). Before submission, the site should be published with public viewing enabled and this link replaced with the live URL.'); set_font(r)

add_para(doc, 'Kitsune Sushi is a five-page restaurant website for Sushi Chat in L’Union, Occitanie. The Home, Menu, Our Story, Gallery, and Order and Reservations pages work together to move a visitor from first impression to a practical action: explore dishes, add items to a cart, place a pickup or delivery request, or reserve a table. The website uses a lively Japanese-inspired visual identity while keeping the core restaurant information—food, location, dining options, and reservation path—easy to find. Its strongest marketing role is to turn interest into a visit or order without making the experience feel like a generic online menu.')

add_heading(doc, 'Marketing Objectives')
add_para(doc, 'The primary objective is conversion: encourage local customers to order food or make a reservation. Prominent “Order Now,” “Explore the Menu,” and “Book Your Table” calls to action appear throughout the site, reducing the distance between browsing and a purchase decision. A secondary objective is brand building. Sushi Chat is positioned as a warm, contemporary destination for Japanese-inspired dining rather than simply a low-cost takeaway option. The website also supports relationship marketing through an email sign-up for the “Sushi Chat Club,” which can be used to share offers, new menu items, and event announcements with interested guests.')

add_heading(doc, 'Value for the Target Audience')
add_para(doc, 'The target audience is local diners and nearby students, workers, families, and social groups who want a casual Japanese meal, either in the restaurant or at home. The site provides value by answering the questions that influence restaurant choice: What is available? How much does it cost? Is there something vegetarian or spicy? Can I dine in, reserve, or order takeaway? Menu filters, a search field, category labels, prices in euros, dish descriptions, vegetarian and hot markers, and a visible cart make this information fast to scan. The visit section gives the L’Union address and reinforces that both takeaway and dine-in are available. Customer testimonials further reduce perceived risk by showing social proof about freshness, service, cleanliness, portions, accessibility, and the atmosphere.')

add_heading(doc, 'Website Content')
add_para(doc, 'Content is organized to suit both exploratory and task-focused visitors. The Home page introduces the brand promise—“A Little Magic. A Lot of Umami.”—then shows signature dishes, a short food-preference quiz, an evening promotion, the brand story, and customer reviews. The Menu page is the practical core of the site: dishes are grouped into categories including nigiri, sashimi, classic maki, signature rolls, hot dishes, bento, veggie choices, and desserts. The Our Story page explains the restaurant’s Japanese inspiration, craft process, team philosophy, and promise of fresh ingredients and warm service. The Gallery page provides visual evidence of food, the restaurant, and behind-the-scenes preparation. Finally, the Order and Reservations page turns content into action through order and reservation forms, cart controls, visit details, and frequently asked questions. This sequence provides useful information before asking for a commitment.')

add_heading(doc, 'Website Design Elements')
add_para(doc, 'The design uses an off-white background, ink-black sections, vivid red accents, large condensed display lettering, playful angled elements, Japanese characters, and a kitsune fox mascot. Together, these elements create a recognizable personality that suits a youthful, bold food brand. Strong image-led cards make menu items appetizing, while visual hierarchy guides attention from the page headline to supporting copy and buttons. Repeated navigation, consistent button shapes, category chips, and the persistent cart create a coherent user experience across all five pages. The contrast between light and dark sections helps segment long pages, and the red accent consistently signals interactive or important elements. Responsive navigation and labeled controls also make the layout more usable on smaller screens.')
add_para(doc, 'The design also supports the restaurant’s storytelling objective. Rather than relying only on polished food photography, it combines food images with the fox character, handwritten-style annotations, “Tokyo” night imagery, and craft-oriented language. This gives visitors an emotional reason to remember the restaurant and creates material that can be reused on social channels. The visual system nevertheless retains practical cues: prices are isolated in a contrasting color, category chips behave consistently, and the most important buttons are repeated at moments of high interest. The result is a site that can entertain a visitor while still helping that visitor decide what to eat and how to order it.')

add_heading(doc, 'Additional Considerations')
add_para(doc, 'First, the site should undergo usability and conversion testing. A short moderated test with local diners could ask users to find a vegetarian meal, reserve a table, and complete a pickup order. Analytics should then measure menu-to-cart rate, completed reservation rate, search use, mobile drop-off, and the call-to-action that produces the most orders. A/B testing could compare button wording such as “Order Now” and “Start Your Order,” or test whether testimonials placed nearer to the reservation form increase bookings.')
add_para(doc, 'The test should deliberately include first-time visitors as well as returning customers. First-time visitors reveal whether the brand story, pricing, and location information are clear without prior knowledge; returning customers reveal whether the menu and order flow are fast enough to support repeat purchases. Observing where users pause, ask questions, or abandon an action is more useful than relying on a general opinion that the page “looks good.” Findings should become an ordered improvement list, beginning with failures that stop an order or reservation and then addressing opportunities that make the experience more persuasive.')
add_para(doc, 'Second, content governance is essential for a restaurant. Prices, dish availability, allergens, opening hours, location details, delivery areas, and promotions can change quickly. The existing note encouraging guests to confirm menu details and allergens is useful, but a production version should connect the menu to an owner-managed source of truth. It should show clear allergen information and a timestamp or “last updated” indicator. This protects customer trust and prevents a successful marketing website from creating disappointment at the point of purchase.')
add_para(doc, 'Third, the website needs technical quality assurance. Testing should cover current versions of Chrome, Safari, Firefox, and Edge; touch use on smaller phones; form validation; keyboard-only navigation; cart quantity changes; empty search results; links; and error or confirmation messages. Performance testing should focus on image compression, lazy loading, and mobile network speed because the site uses many large food and atmosphere photographs. A privacy review is also needed before collecting newsletter or order information: forms should explain how data will be used and be sent over HTTPS.')

add_heading(doc, 'Recommended Enhancement Using an Outside Source')
add_para(doc, 'The most valuable enhancement would be a formal accessibility pass targeting WCAG 2.2 AA. The World Wide Web Consortium explains that WCAG 2.2 organizes accessibility guidance under four principles—perceivable, operable, understandable, and robust—and provides testable success criteria. Applied here, that means checking color contrast for red, cream, and gray text; making every image alt text meaningful; ensuring visible keyboard focus; keeping navigation and modal gallery controls fully operable without a mouse; associating form labels and error messages correctly; and giving buttons unambiguous names. The current website already includes descriptive image alt text and labels on many controls, which is a good foundation. A systematic WCAG review would make the colorful experience usable for more guests, including people using screen readers, keyboards, or mobile accessibility settings, while reducing friction in the high-value ordering and reservation journey.')

add_heading(doc, 'Reference')
p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1; p.paragraph_format.space_after = Pt(0); p.paragraph_format.first_line_indent = Inches(0)
r = p.add_run('World Wide Web Consortium. “WCAG 2 Overview.” Web Accessibility Initiative, 17 Sept. 2026, '); set_font(r)
add_link(p, 'https://www.w3.org/WAI/standards-guidelines/wcag/', 'https://www.w3.org/WAI/standards-guidelines/wcag/')
r = p.add_run('. Accessed 25 Oct. 2026.'); set_font(r)

doc.save(OUT)
print(OUT)
